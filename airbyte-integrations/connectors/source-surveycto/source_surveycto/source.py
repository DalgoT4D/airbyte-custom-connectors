#
# Copyright (c) 2025 Airbyte, Inc., all rights reserved.
#


import base64
import io
import re
import time
from abc import ABC
from typing import Any, Iterable, List, Mapping, MutableMapping, Optional, Tuple

import openpyxl
import requests
from airbyte_cdk.models import SyncMode
from airbyte_cdk.sources import AbstractSource
from airbyte_cdk.sources.streams import CheckpointMixin, Stream
from airbyte_cdk.sources.streams.http import HttpStream
from airbyte_cdk.sources.utils.transform import TransformConfig, TypeTransformer


stream_json_schema = {
    "$schema": "http://json-schema.org/draft-07/schema#",
    "type": "object",
    "additionalProperties": True,
    "properties": {
        "KEY": {
            "type": [
                "string",
                "null",
            ]
        },
        "endtime": {"type": ["string", "null"]},
        "data": {
            "type": "object",
        },
        "SubmissionDate": {"type": ["string", "null"]},
    },
}

form_definitions_json_schema = {
    "$schema": "http://json-schema.org/draft-07/schema#",
    "type": "object",
    "additionalProperties": True,
    "properties": {
        "form_id": {"type": ["string", "null"]},
        "form_version": {"type": ["string", "null"]},
        "sheet": {"type": ["string", "null"]},
        "row_index": {"type": ["number", "null"]},
        "row": {"type": "object"},
    },
}

# same key-cleaning as agency-fund/rtv-questionsdb applied to XLSForm headers
FORM_DEFINITION_KEY_CLEAN_RE = re.compile(r"[ ,;{}()\n\t=-]")

# the three standard XLSForm sheets
XLSFORM_SHEETS = ("survey", "choices", "settings")


class SurveyStream(HttpStream, ABC):
    transformer: TypeTransformer = TypeTransformer(TransformConfig.DefaultSchemaNormalization)

    def __init__(self, config: Mapping[str, Any], form_id, schema, **kwargs):
        self.form_id = None
        super().__init__()

        self.config = config
        self.schema = schema
        self.server_name = config["server_name"]
        self.form_id = form_id
        self.start_date = config["start_date"]
        # base64 encode username and password as auth token
        user_name_password = f"{config['username']}:{config['password']}"
        self.auth_token = self._base64_encode(user_name_password)

    @property
    def url_base(self) -> str:
        return f"https://{self.server_name}.surveycto.com/" + "api/v2/forms/data/wide/json/"

    def _base64_encode(self, string: str) -> str:
        return base64.b64encode(string.encode("ascii")).decode("ascii")

    def request_params(
        self, stream_state: Mapping[str, Any], stream_slice: Mapping[str, any] = None, next_page_token: Mapping[str, Any] = None
    ) -> MutableMapping[str, Any]:
        return {}


class SurveyctoStream(SurveyStream, CheckpointMixin):
    primary_key = "KEY"
    cursor_field = "SubmissionDate"
    _cursor_value = None

    @property
    def state(self) -> Mapping[str, Any]:
        if self._cursor_value:
            return {self.cursor_field: self._cursor_value}
        else:
            return {self.cursor_field: self.start_date}

    @state.setter
    def state(self, value: Mapping[str, Any]):
        if self.cursor_field in value:
            self._cursor_value = value[self.cursor_field]

    @property
    def name(self) -> str:
        return getattr(self, 'form_id', 'Surveyctostream')

    def next_page_token(self, response: requests.Response) -> Optional[Mapping[str, Any]]:
        return None

    def get_json_schema(self):
        return self.schema

    def request_params(
        self, stream_state: Mapping[str, Any], stream_slice: Mapping[str, any] = None, next_page_token: Mapping[str, Any] = None
    ) -> MutableMapping[str, Any]:
        ix = self.state[self.cursor_field]
        return {"date": ix}

    def request_headers(
        self, stream_state: Mapping[str, Any], stream_slice: Mapping[str, Any] = None, next_page_token: Mapping[str, Any] = None
    ) -> Mapping[str, Any]:
        return {"Authorization": "Basic " + self.auth_token}

    def path(self, stream_slice: Mapping[str, Any] = None, **kwargs) -> str:
        return self.form_id

    def parse_response(
        self,
        response: requests.Response,
        stream_state: Mapping[str, Any],
        stream_slice: Mapping[str, Any] = None,
        next_page_token: Mapping[str, Any] = None,
    ) -> Iterable[Mapping]:
        self.response_json = response.json()

        for record in self.response_json:
            # send data, key, submission date and endtime
            record_id = record.get("KEY")
            submission_date = record.get("SubmissionDate")
            endtime = record.get("endtime")

            retval = {"KEY": record_id, "data": record}
            retval["SubmissionDate"] = submission_date
            retval["endtime"] = endtime

            yield retval

    def read_records(self, *args, **kwargs) -> Iterable[Mapping[str, Any]]:
        for record in super().read_records(*args, **kwargs):
            self._cursor_value = record[self.cursor_field]
            yield record


class SurveyctoFormDefinitions(HttpStream):
    """
    Form-definition stream: one stream (`form_definitions`) covering every form
    in the config. For each form_id we GET /forms/{form_id}/files, which lists
    the currently deployed XLSForm definition file (deployedGroupFiles.definitionFile)
    and all previous definition files (previousDefinitionFiles). We download each
    definition .xlsx and emit one record per row of its survey/choices/settings
    sheets, so downstream models (e.g. dbt_taf_metrics QuestionsDB) can join
    question labels/types/choices to melted submissions.

    Historical versions are included on purpose: submissions carry a
    formdef_version field stamped with the form version that was live at
    submission time, and the join must find that version.

    Full-refresh only; meant to be synced with destination mode "overwrite".
    Definition files are small, so re-downloading everything each sync is fine.
    """

    primary_key = ["form_id", "form_version", "sheet", "row_index"]

    # full-refresh stream: no cursor field
    cursor_field = []

    def __init__(self, config: Mapping[str, Any], form_ids: List[str], **kwargs):
        """constructor"""
        super().__init__()
        self.config = config
        self.server_name = config["server_name"]
        self.form_ids = form_ids
        user_name_password = f"{config['username']}:{config['password']}"
        self.auth_token = base64.b64encode(user_name_password.encode("ascii")).decode("ascii")

    @property
    def url_base(self) -> str:
        """base url for all http requests for this server"""
        return f"https://{self.server_name}.surveycto.com/"

    @property
    def name(self) -> str:
        """single stream for all forms"""
        return "form_definitions"

    def get_json_schema(self):
        """airbyte needs this function"""
        return form_definitions_json_schema

    def stream_slices(
        self,
        sync_mode: SyncMode = None,  # pylint:disable=unused-argument
        cursor_field: List[str] = None,  # pylint:disable=unused-argument
        stream_state: Mapping[str, Any] = None,  # pylint:disable=unused-argument
    ) -> Iterable[Optional[Mapping[str, Any]]]:
        """one slice per form"""
        for form_id in self.form_ids:
            yield {"form_id": form_id}

    def path(self, stream_slice: Mapping[str, Any] = None, **kwargs) -> str:  # pylint:disable=unused-argument
        """files-metadata endpoint for one form"""
        return f"forms/{stream_slice['form_id']}/files"

    def next_page_token(self, response: requests.Response) -> Optional[Mapping[str, Any]]:  # pylint:disable=unused-argument
        """the files endpoint is not paginated"""
        return None

    def request_params(
        self, stream_state: Mapping[str, Any], stream_slice: Mapping[str, Any] = None, next_page_token: Mapping[str, Any] = None
    ) -> MutableMapping[str, Any]:
        """the files endpoint requires a unix-milliseconds timestamp param"""
        return {"t": int(time.time() * 1000)}

    def request_headers(
        self,
        stream_state: Mapping[str, Any],  # pylint:disable=unused-argument
        stream_slice: Mapping[str, Any] = None,  # pylint:disable=unused-argument
        next_page_token: Mapping[str, Any] = None,  # pylint:disable=unused-argument
    ) -> Mapping[str, Any]:
        """build the request headers"""
        return {"Authorization": "Basic " + self.auth_token}

    def parse_response(
        self,
        response: requests.Response,
        stream_state: Mapping[str, Any] = None,  # pylint:disable=unused-argument
        stream_slice: Mapping[str, Any] = None,
        next_page_token: Mapping[str, Any] = None,  # pylint:disable=unused-argument
    ) -> Iterable[Mapping]:
        """download every definition-file version listed for this form and
        yield one record per XLSForm sheet row"""
        form_id = stream_slice["form_id"] if stream_slice else None
        files_json = response.json() or {}

        definition_files = []
        deployed_group_files = files_json.get("deployedGroupFiles") or {}
        if deployed_group_files.get("definitionFile"):
            definition_files.append(deployed_group_files["definitionFile"])
        # previous versions are included on purpose: submissions'
        # formdef_version must join to the version live at submission time
        previous_files = files_json.get("previousDefinitionFiles") or deployed_group_files.get("previousDefinitionFiles") or []
        definition_files.extend(previous_files)

        seen_versions = set()
        for definition_file in definition_files:
            form_version = definition_file.get("formVersion")
            form_version = str(form_version) if form_version is not None else None
            if form_version in seen_versions:
                continue
            seen_versions.add(form_version)
            yield from self._parse_definition_file(
                form_id=form_id,
                form_version=form_version,
                filename=definition_file.get("filename") or "",
                download_link=definition_file.get("downloadLink"),
            )

    def _parse_definition_file(
        self, form_id: str, form_version: str, filename: str, download_link: str
    ) -> Iterable[Mapping[str, Any]]:
        """download one XLSForm definition file and yield one record per row of
        its survey/choices/settings sheets; skip (with a warning) rather than
        fail the sync when a single file cannot be fetched or parsed"""
        if not download_link:
            self.logger.warning("form %s version %s: definition file has no downloadLink, skipping", form_id, form_version)
            return
        if filename.lower().endswith(".xls"):
            self.logger.warning(
                "form %s version %s: legacy .xls definition file %s is not supported, skipping", form_id, form_version, filename
            )
            return

        try:
            file_response = requests.get(download_link, headers={"Authorization": "Basic " + self.auth_token}, timeout=300)
            file_response.raise_for_status()
        except requests.RequestException as error:
            self.logger.warning("form %s version %s: failed to download definition file %s (%s), skipping", form_id, form_version, filename, error)
            return

        try:
            workbook = openpyxl.load_workbook(io.BytesIO(file_response.content), read_only=True, data_only=True)
        except Exception as error:  # pylint:disable=broad-except
            self.logger.warning("form %s version %s: failed to parse definition file %s (%s), skipping", form_id, form_version, filename, error)
            return

        try:
            sheet_names = {sheet_name.strip().lower(): sheet_name for sheet_name in workbook.sheetnames}
            for sheet in XLSFORM_SHEETS:
                if sheet not in sheet_names:
                    continue
                worksheet = workbook[sheet_names[sheet]]
                rows = worksheet.iter_rows(values_only=True)
                header = next(rows, None)
                if header is None:
                    continue
                columns = [self._clean_column_name(column) for column in header]
                # row_index is the 0-based position of the row under the header
                # row; question order matters downstream, and blank rows keep
                # their slot so the ordering is stable
                for row_index, values in enumerate(rows):
                    row = {
                        column: str(value)
                        for column, value in zip(columns, values)
                        if column is not None and value is not None
                    }
                    if not row:
                        continue
                    yield {
                        "form_id": form_id,
                        "form_version": form_version,
                        "sheet": sheet,
                        "row_index": row_index,
                        "row": row,
                    }
        finally:
            workbook.close()

    @staticmethod
    def _clean_column_name(column: Any) -> Optional[str]:
        """clean an XLSForm header cell the same way rtv-questionsdb did:
        replace [ ,;{}()\\n\\t=-] with '_', lowercase, and drop unnamed columns"""
        if column is None:
            return None
        cleaned = FORM_DEFINITION_KEY_CLEAN_RE.sub("_", str(column).strip()).lower()
        if not cleaned or cleaned.startswith("unnamed"):
            return None
        return cleaned


# Source
class SourceSurveycto(AbstractSource):
    def check_connection(self, logger, config) -> Tuple[bool, Any]:

        form_ids = config["form_id"]

        try:
            for form_id in form_ids:
                stream = SurveyctoStream(config=config, form_id=form_id, schema=stream_json_schema)
                next(stream.read_records(sync_mode=SyncMode.full_refresh))

            return True, None

        except Exception as error:
            return False, f"Unable to connect - {(error)}"

    def generate_streams(self, config: str) -> List[Stream]:
        forms = config.get("form_id", [])
        streams = []

        for form_id in forms:
            stream = SurveyctoStream(config=config, form_id=form_id, schema=stream_json_schema)
            streams.append(stream)

        # Single form-definition stream covering all configured forms, used by
        # dbt_taf_metrics (QuestionsDB) to join question labels/types/choices
        # to melted submissions
        if forms:
            streams.append(SurveyctoFormDefinitions(config=config, form_ids=forms))

        return streams

    def streams(self, config: Mapping[str, Any]) -> List[Stream]:
        streams = self.generate_streams(config=config)
        return streams
